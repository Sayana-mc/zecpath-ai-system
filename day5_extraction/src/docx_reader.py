from docx import Document


def extract_text_from_docx(file_path):
    """
    Extract text from a DOCX resume.
    """

    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Extract text from tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                paragraphs.append(" | ".join(row_text))

    return "\n".join(paragraphs)